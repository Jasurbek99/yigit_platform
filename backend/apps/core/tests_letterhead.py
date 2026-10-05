"""Letterhead helpers: find/fill the «№ ___» blank, validate uploads (spec 2026-10-05)."""
import shutil
import tempfile
from io import BytesIO

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase, TestCase, override_settings
from docx import Document
from rest_framework.test import APIClient

from apps.core.letterhead import fill_number, number_blank_run, validate_letterhead
from apps.core.models import ExportFirm, User


def letterhead_bytes(*runs: str) -> bytes:
    """A docx whose second body paragraph is made of the given runs."""
    doc = Document()
    doc.add_paragraph('ÝIGIT HOJALYK JEMGYÝETI')
    para = doc.add_paragraph()
    for text in runs:
        para.add_run(text)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _doc(data: bytes):
    return Document(BytesIO(data))


class NumberBlankTest(SimpleTestCase):
    def test_finds_blank_in_one_run(self):
        doc = _doc(letterhead_bytes('«____» ______ 20 ý.', '    № ___________ '))
        self.assertIn('___', number_blank_run(doc).text)

    def test_number_split_across_runs(self):
        doc = _doc(letterhead_bytes('№', ' ', '__________'))
        fill_number(doc, 15)
        self.assertIn('№ 15', doc.paragraphs[1].text)

    def test_date_underscores_before_number_sign_are_left(self):
        doc = _doc(letterhead_bytes('«____» ______ 20 ý.', '    № ___________ '))
        fill_number(doc, 7)
        text = doc.paragraphs[1].text
        self.assertTrue(text.startswith('«____» ______ 20 ý.'))
        self.assertIn('№ 7', text)

    def test_date_and_number_in_one_run(self):
        doc = _doc(letterhead_bytes('«____» ______ 20 ý.    № ________'))
        fill_number(doc, 3)
        self.assertEqual(doc.paragraphs[1].text, '«____» ______ 20 ý.    № 3')

    def test_underscores_split_across_runs(self):
        # Word splits a blank into short runs on any formatting change.
        doc = _doc(letterhead_bytes('№ __', '__', '__'))
        self.assertIsNotNone(number_blank_run(doc))
        fill_number(doc, 7)
        self.assertEqual(doc.paragraphs[1].text, '№ 7')

    def test_blank_inside_a_table_cell(self):
        doc = Document()
        doc.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0].add_run('№ _____')
        fill_number(doc, 9)
        self.assertEqual(doc.tables[0].cell(0, 0).text, '№ 9')

    def test_no_number_sign_means_no_blank(self):
        self.assertIsNone(number_blank_run(_doc(letterhead_bytes('«____» ______ 20 ý.'))))

    def test_none_number_leaves_the_blank(self):
        doc = _doc(letterhead_bytes('№ _____'))
        fill_number(doc, None)
        self.assertIn('№ _____', doc.paragraphs[1].text)


class ValidateLetterheadTest(SimpleTestCase):
    def test_valid_file_passes(self):
        validate_letterhead(BytesIO(letterhead_bytes('№ _____')))

    def test_missing_blank_is_rejected(self):
        with self.assertRaisesMessage(ValueError, '№'):
            validate_letterhead(BytesIO(letterhead_bytes('no blank here')))

    def test_content_in_page_header_is_rejected(self):
        # docxcompose drops Word's page header, so a logo there would never print.
        doc = Document(BytesIO(letterhead_bytes('№ ___')))
        doc.sections[0].header.paragraphs[0].add_run('ÝIGIT')
        buf = BytesIO()
        doc.save(buf)
        with self.assertRaises(ValueError):
            validate_letterhead(BytesIO(buf.getvalue()))

    def test_corrupt_docx_is_rejected(self):
        with self.assertRaises(ValueError):
            validate_letterhead(BytesIO(b'PK\x03\x04 not really a zip'))

    def test_truncated_docx_is_rejected(self):
        with self.assertRaises(ValueError):
            validate_letterhead(BytesIO(letterhead_bytes('№ ___')[:4000]))


class LetterheadUploadApiTest(TestCase):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls._media = tempfile.mkdtemp()
        cls._override = override_settings(MEDIA_ROOT=cls._media)
        cls._override.enable()

    @classmethod
    def tearDownClass(cls):
        cls._override.disable()
        shutil.rmtree(cls._media, ignore_errors=True)
        super().tearDownClass()

    def setUp(self):
        self.firm = ExportFirm.objects.create(code='LHUP', name_tk='LH', name_short='LH')
        self.client = APIClient()
        self.client.force_authenticate(User.objects.create_user(
            username='lh_admin', password='x', role='admin', is_superuser=True,
        ))
        self.url = f'/api/v1/export/admin/firms/{self.firm.id}/'

    def _upload(self, data: bytes, name='blank.docx'):
        return self.client.patch(
            self.url, {'letterhead': SimpleUploadedFile(name, data)}, format='multipart',
        )

    def test_valid_upload_is_saved(self):
        resp = self._upload(letterhead_bytes('№ _____'))
        self.assertEqual(resp.status_code, 200, resp.content)
        self.assertTrue(resp.json()['letterhead'].startswith('/media/export_firms/letterheads/'))

    def test_bad_upload_is_400(self):
        resp = self._upload(letterhead_bytes('nothing'))
        self.assertEqual(resp.status_code, 400)
        self.assertIn('letterhead', resp.json())

    def test_non_docx_name_is_400(self):
        self.assertEqual(self._upload(letterhead_bytes('№ ___'), name='blank.pdf').status_code, 400)
