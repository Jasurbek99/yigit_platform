"""CT-1 / Fito / ARZA on the seller's letterhead (spec 2026-10-05 §5)."""
import shutil
import tempfile
import unittest
from decimal import Decimal
from io import BytesIO

from django.core.files.base import ContentFile
from django.test import TestCase, override_settings
from docx import Document
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.shared import Mm, Pt

from apps.contracts.services import document_render as render
from apps.contracts.services.document_render import generate
from apps.contracts.services.letterhead_render import apply_letterhead
from apps.contracts.tests.test_contract_sale_api import (
    _make_contract, _make_export_firm, _make_import_firm, _make_invoice, _make_season,
)
from apps.contracts.tests.test_document_generation import _make_packed_shipment
from apps.core.tests_letterhead import letterhead_bytes

LETTER_KEYS = ('ct1_ru', 'fito_ru', 'customs_tk')


def _text(docx_bytes: bytes) -> str:
    return '\n'.join(p.text for p in Document(BytesIO(docx_bytes)).paragraphs)


def _save(doc) -> bytes:
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _letter(*, line_spacing=None, trailing_empty=False) -> bytes:
    doc = Document()
    if line_spacing is not None:
        doc.styles['Normal'].paragraph_format.line_spacing = line_spacing
    doc.add_paragraph('Гос служба по карантину')
    if trailing_empty:
        doc.add_paragraph('')
    return _save(doc)


def _two_section_letterhead() -> bytes:
    """Like the real Yigit blank: its own section (columns), then a continuous break."""
    doc = Document(BytesIO(letterhead_bytes('ÝIGIT', ' № ______')))
    doc.add_section(WD_SECTION.CONTINUOUS)
    return _save(doc)


class ApplyLetterheadTest(TestCase):
    def test_letterhead_goes_on_top_with_number(self):
        out = apply_letterhead(_letter(), letterhead_bytes('ÝIGIT', ' № ______'), 15)
        text = _text(out)
        self.assertLess(text.index('ÝIGIT'), text.index('Гос служба'))
        self.assertIn('№ 15', text)

    def test_no_number_keeps_the_blank(self):
        out = apply_letterhead(_letter(), letterhead_bytes('№ ______'), None)
        self.assertIn('№ ______', _text(out))

    def test_letter_keeps_its_last_section_margins(self):
        letter = Document(BytesIO(_letter()))
        letter.sections[-1].left_margin = Pt(77)
        out = Document(BytesIO(apply_letterhead(_save(letter), _two_section_letterhead(), 1)))
        self.assertEqual(out.sections[-1].left_margin, Pt(77))

    def test_letterhead_takes_the_letters_page_size(self):
        letter = Document(BytesIO(_letter()))
        letter.sections[-1].page_width, letter.sections[-1].page_height = Mm(210), Mm(297)
        head = Document(BytesIO(_two_section_letterhead()))
        for section in head.sections:  # a US Letter blank
            section.page_width, section.page_height = Mm(216), Mm(279)
        out = Document(BytesIO(apply_letterhead(_save(letter), _save(head), 1)))
        sizes = {(s.page_width, s.page_height) for s in out.sections}
        self.assertEqual(sizes, {(out.sections[-1].page_width, out.sections[-1].page_height)})
        self.assertAlmostEqual(out.sections[0].page_height.mm, 297, delta=0.1)

    def test_letter_section_is_continuous(self):
        out = Document(BytesIO(apply_letterhead(_letter(), _two_section_letterhead(), 1)))
        self.assertGreater(len(out.sections), 1)
        self.assertEqual(out.sections[-1].start_type, WD_SECTION.CONTINUOUS)

    def test_spacing_frozen_per_attribute(self):
        head = Document(BytesIO(letterhead_bytes('ÝIGIT', ' № ______')))
        head.styles['Normal'].paragraph_format.space_after = Pt(0)  # sets only w:after
        out = Document(BytesIO(apply_letterhead(_letter(line_spacing=1.5), _save(head), 1)))
        para = next(p for p in out.paragraphs if 'ÝIGIT' in p.text)
        spacing = para._p.pPr.find(qn('w:spacing'))
        self.assertEqual(spacing.get(qn('w:after')), '0')
        self.assertNotEqual(spacing.get(qn('w:line')), '360')  # not the letter's 1.5 lines

    def test_style_font_survives_merge(self):
        head = Document(BytesIO(letterhead_bytes('ÝIGIT', ' № ______')))
        normal = head.styles['Normal']
        normal.font.name = 'Arial'
        normal.font.size = Pt(8)
        letter = Document(BytesIO(_letter()))
        letter.styles['Normal'].font.name = 'Times New Roman'
        letter.styles['Normal'].font.size = Pt(14)
        out = Document(BytesIO(apply_letterhead(_save(letter), _save(head), 1)))
        run = next(p for p in out.paragraphs if 'ÝIGIT' in p.text).runs[0]
        self.assertEqual(run.font.name, 'Arial')
        self.assertEqual(run.font.size, Pt(8))
        letter_run = next(p for p in out.paragraphs if 'Гос служба' in p.text).runs[0]
        self.assertIsNone(letter_run.font.name)  # the letter still follows its own Normal

    def test_letterheads_trailing_empty_lines_trimmed(self):
        head = Document(BytesIO(letterhead_bytes('ÝIGIT', ' № ______')))
        for _ in range(3):
            head.add_paragraph('')
        out = Document(BytesIO(apply_letterhead(_letter(), _save(head), 1)))
        texts = [p.text for p in out.paragraphs]
        number_line = next(i for i, t in enumerate(texts) if '№ 1' in t)
        self.assertEqual(texts[number_line + 1], 'Гос служба по карантину')

    def test_letter_body_is_left_alone(self):
        out = Document(BytesIO(apply_letterhead(_letter(trailing_empty=True), letterhead_bytes('№ ___'), 1)))
        self.assertEqual(out.paragraphs[-1].text, '')

    def test_paragraph_after_letterhead_table_kept(self):
        head = Document(BytesIO(letterhead_bytes('№ ___')))
        head.add_table(rows=1, cols=1)
        head.add_paragraph('')
        out = Document(BytesIO(apply_letterhead(_letter(), _save(head), 1)))
        body = [e.tag for e in out.element.body]
        table_at = body.index(qn('w:tbl'))
        self.assertEqual(body[table_at + 1], qn('w:p'))
        self.assertEqual(''.join(out.element.body[table_at + 1].itertext()), '')


class GenerateWithLetterheadTest(TestCase):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls._media = tempfile.mkdtemp()
        cls._override = override_settings(MEDIA_ROOT=cls._media)
        cls._override.enable()

    @classmethod
    def tearDownClass(cls):
        cls._override.disable()
        shutil.rmtree(cls._media, ignore_errors=True)
        super().tearDownClass()

    def setUp(self):
        season = _make_season()
        imp = _make_import_firm('IMPLHR')
        self.firm = _make_export_firm('LHRFIRM')
        contract = _make_contract('LHR-1', self.firm, imp, season)
        self.sale = _make_invoice(contract, invoice_number=1)
        self.sale.shipment = _make_packed_shipment(season, imp, code='0707001/25')
        self.sale.export_firm = self.firm
        self.sale.save(update_fields=['shipment', 'export_firm'])
        self.sale.refresh_from_db()  # dates given as 'YYYY-MM-DD' stay str in memory
        self.sale.shipment.refresh_from_db()

    def _attach(self, firm, data: bytes):
        firm.letterhead.save('lh.docx', ContentFile(data))

    def test_no_letterhead_is_unchanged(self):
        data, _, _ = generate('ct1_ru', self.sale, 'docx', {}, highlight=False)
        self.assertNotIn('ÝIGIT', _text(data))
        self.assertIn('СТ-1', _text(data))

    def test_with_letterhead_and_number(self):
        self._attach(self.firm, letterhead_bytes('ÝIGIT', '№ ____'))
        self.sale.ct1_number = 7
        self.sale.save(update_fields=['ct1_number'])
        data, _, _ = generate('ct1_ru', self.sale, 'docx', {}, highlight=False)
        self.assertIn('№ 7', _text(data))
        self.assertIn('ÝIGIT', _text(data))

    def test_each_letter_prints_its_own_number(self):
        self._attach(self.firm, letterhead_bytes('ÝIGIT', '№ ____'))
        self.sale.ct1_number, self.sale.fito_number, self.sale.customs_number = 3, 4, 5
        self.sale.save(update_fields=['ct1_number', 'fito_number', 'customs_number'])
        for key, number in zip(LETTER_KEYS, (3, 4, 5)):
            data, _, _ = generate(key, self.sale, 'docx', {}, highlight=False)
            self.assertIn(f'№ {number}', _text(data), key)

    def test_other_firms_letterhead_never_used(self):
        self._attach(_make_export_firm('OTHERLH'), letterhead_bytes('OTHERHEAD', '№ ____'))
        data, _, _ = generate('fito_ru', self.sale, 'docx', {}, highlight=False)
        self.assertNotIn('OTHERHEAD', _text(data))

    def test_invoice_never_gets_letterhead(self):
        self._attach(self.firm, letterhead_bytes('ÝIGIT', '№ ____'))
        data, _, _ = generate('invoice_ru', self.sale, 'docx', {'place_loading': 'Kaka'}, highlight=False)
        self.assertNotIn('ÝIGIT', _text(data))

    def test_unreadable_letterhead_renders_plain(self):
        self._attach(self.firm, b'broken')
        with self.assertLogs('apps.contracts.services.letterhead_render', 'WARNING'):
            data, _, _ = generate('customs_tk', self.sale, 'docx', {}, highlight=False)
        self.assertIn('ARZA', _text(data))

    def test_merge_failure_renders_plain(self):
        from unittest import mock
        self._attach(self.firm, letterhead_bytes('ÝIGIT', '№ ____'))
        with mock.patch('apps.contracts.services.letterhead_render.apply_letterhead',
                        side_effect=KeyError('missing style')):
            with self.assertLogs('apps.contracts.services.document_render', 'ERROR'):
                data, _, _ = generate('fito_ru', self.sale, 'docx', {}, highlight=False)
        self.assertNotIn('ÝIGIT', _text(data))
        self.assertIn('Фитосанитарный', _text(data))

    def test_packet_letters_carry_the_letterhead(self):
        import zipfile
        self._attach(self.firm, letterhead_bytes('ÝIGIT', '№ ____'))
        packet = render.generate_packet_zip(
            self.sale.shipment, [self.sale], 'ru', 'docx', {'place_loading': 'Kaka'}, False,
        )
        archive = zipfile.ZipFile(BytesIO(packet))
        letters = [n for n in archive.namelist() if n.startswith(('CT1_', 'Fito_', 'Customs_'))]
        self.assertEqual(len(letters), 3, archive.namelist())
        for name in letters:
            self.assertIn('ÝIGIT', _text(archive.read(name)), name)

    @unittest.skipUnless(render._libreoffice_bin(), 'LibreOffice not installed')
    def test_pdf_one_page_each(self):
        import fitz

        # A blank as compact as the real Yigit one (top margin ~6mm). A taller blank
        # can legitimately push ARZA to a 2nd page — that is the blank's design.
        head = Document(BytesIO(_two_section_letterhead()))
        head.styles['Normal'].paragraph_format.space_after = Pt(0)
        head.sections[0].top_margin = Mm(6)
        self._attach(self.firm, _save(head))
        for key in LETTER_KEYS:
            data, _, _ = generate(key, self.sale, 'pdf', {}, highlight=False)
            self.assertEqual(fitz.open(stream=data, filetype='pdf').page_count, 1, key)
