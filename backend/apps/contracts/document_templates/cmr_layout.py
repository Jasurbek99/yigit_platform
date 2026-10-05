"""Where each CMR overlay value prints on the office's blank CMR form.

Positions were measured off a 200-dpi scan of the blank form (``data/cmr.PDF``).
The paper is about 204 × 290 mm, a little smaller than A4. Every number here is in
**millimetres from the paper's top-left corner**, so the layout is the paper's
own geometry and does not depend on any spreadsheet's column units.

``FIELDS`` is the single source of truth. ``build_cmr_docx`` moves the Word
form's frames onto it (the Word CMR is what gets printed on the blank — its
positions are absolute). ``build_cmr_xlsx`` turns it into a spreadsheet grid, and
``document_context.build_cmr_overlay`` uses ``FIELD_CELLS`` to find the cell for
each value, so those two cannot drift apart. RU and EN share
one layout because the physical form is the same; only the fixed labels differ.

Writing lines on the form are ~4.8 mm apart. A single-line field names the line
its text sits on (``first``); a wrapped field covers ``first`` .. ``last``.

Registration: the sheet prints on A4 at 100% with the grid starting at the
left/top margins (``ORIGIN_MM``). If a test print is shifted on the blank, the
office corrects it at print time by changing those two margins in Excel — every
field moves together.
"""
from typing import NamedTuple

from openpyxl.utils import get_column_letter

# Grid origin on the paper = the sheet's left and top print margins. Kept clear of
# the printer's unprintable edge; the first field starts well inside it.
ORIGIN_MM = 10.0

# How far a field's cell reaches above / below the writing line it sits on.
ABOVE_LINE_MM = 3.5
BELOW_LINE_MM = 1.1

# Box borders measured on the form (mm from the paper's left edge).
LEFT_EDGE = 14.6      # boxes 1–5 left border
MID = 102.7           # boxes 1–5 | title, 16–18
RIGHT_EDGE = 190.9


class Field(NamedTuple):
    key: str            # overlay-value key, or ``label:<name>`` for a fixed label
    x0: float           # left edge of the text area (mm)
    x1: float           # right edge of the text area (mm)
    first: float        # writing line the (first) text line sits on (mm from top)
    last: float | None = None   # last writing line for a wrapped field
    size: float = 9     # font size, pt
    align: str = 'left'


FIELDS: tuple[Field, ...] = (
    # Box 1 — sender(s). Name after the printed box label, address wrapped below.
    Field('sender1_name', 52.0, 101.5, 25.1),
    Field('sender1_address', 17.0, 101.0, 29.8, 34.5, size=8),
    Field('sender2_name', 17.0, 101.5, 39.2),
    Field('sender2_address', 17.0, 101.0, 43.8, 47.1, size=8),
    # Box 2 — consignee. The address carries its own line breaks (VAT, reg. no…).
    Field('consignee_name', 17.0, 101.5, 58.1),
    Field('consignee_address', 17.0, 101.0, 62.9, 75.8, size=8.5),
    # Box 3 — after the "Döwleti / Country" label.
    Field('country_destination', 37.0, 101.5, 91.9),
    # Box 4 — place (region + etrap side by side), country, date.
    Field('place_region', 32.0, 66.0, 106.5),
    Field('place_district', 66.0, 101.5, 106.5),
    Field('country_dispatch', 40.0, 101.5, 111.2),
    Field('doc_date', 40.0, 80.0, 115.8),
    # Box 5 — annexed documents.
    Field('invoice_refs', 20.0, 101.5, 125.8),
    Field('tir_line', 20.0, 101.5, 130.6),
    # Box 17 — successive carrier, typed at generate-time (optional). Right
    # column, the four writing lines under the box label.
    Field('successive_carrier', 105.0, 190.0, 96.7, 111.3),
    # Boxes 6–12 — cargo. Box 7 = 48.5, 8 = 75.5, 9 = 99.5, 10 = 125.3,
    # 11 = 147.2, 12 = 169.2. Weights go in box 11, net in box 12.
    Field('cargo_name', 100.0, 146.5, 149.8),
    Field('label:gross', 148.0, 168.5, 149.8),
    Field('label:net', 170.0, 190.5, 149.8),
    Field('boxes', 49.0, 75.0, 154.5, align='center'),
    Field('packing', 76.0, 99.0, 154.5),
    Field('label:pallet_weight', 100.0, 146.5, 154.5),
    Field('pallet_weight', 148.0, 160.5, 154.5, align='right'),
    Field('label:kg', 161.0, 168.5, 154.5),
    Field('pallets_line', 20.0, 99.0, 159.3),
    Field('label:gross_without_pallet', 100.0, 146.5, 159.3),
    Field('gross_without_pallet', 148.0, 160.5, 159.3, align='right'),
    Field('label:kg', 161.0, 168.5, 159.3),
    Field('label:gross_with_pallet', 100.0, 146.5, 164.0),
    Field('gross_with_pallet', 148.0, 160.5, 164.0, align='right'),
    Field('label:kg', 161.0, 168.5, 164.0),
    Field('net_line', 170.0, 190.5, 164.0),
    # Box 21 — a 5 mm strip: place after "Resmileşdirilen ýeri", date after "we senesi".
    Field('label:city', 45.0, 84.0, 236.6),
    Field('doc_date', 97.0, 135.0, 236.6),
    # Box 23 — drivers after "Sürüjileriň atlary", passport on the next line.
    Field('driver_name', 94.0, 136.0, 246.7),
    Field('driver_passport', 80.0, 136.0, 251.4),
    # Boxes 25 / 26 — one line per tractor (a second head goes on the next line):
    # model after "Truck", its head/trailer plates level with it in box 26.
    Field('truck_model', 35.0, 66.5, 270.8, 275.7),
    Field('plates', 69.0, 116.0, 270.8, 275.7),
)

# Fixed labels printed with the data (the blank form does not carry them).
LABELS = {
    'ru': {
        'gross': 'Брутто:', 'net': 'Нетто:', 'pallet_weight': 'вес поддона',
        'kg': 'кг.', 'gross_without_pallet': 'вес брутто без подд.',
        'gross_with_pallet': 'вес брутто с подд.', 'city': 'г. Ашгабат',
    },
    'en': {
        'gross': 'GROSS:', 'net': 'NETTO:', 'pallet_weight': 'pallet weight',
        'kg': 'kg.', 'gross_without_pallet': 'gross weight without pallet',
        'gross_with_pallet': 'gross weight with pallet', 'city': 'Ashgabat city',
    },
}


def band(field: Field) -> tuple[float, float]:
    """Top and bottom (mm) of the cell a field prints in."""
    return field.first - ABOVE_LINE_MM, (field.last or field.first) + BELOW_LINE_MM


def _px(mm: float) -> int:
    """Millimetres from the grid origin → whole screen pixels (96 dpi).

    Excel snaps row heights and column widths to whole pixels. Rounding each
    *boundary* (not each size) keeps that rounding from accumulating down the page.
    """
    return round((mm - ORIGIN_MM) / 25.4 * 96)


def _boundaries(values) -> list[int]:
    return sorted({0, *(_px(v) for v in values)})


COL_PX = _boundaries(v for f in FIELDS for v in (f.x0, f.x1))
ROW_PX = _boundaries(v for f in FIELDS for v in band(f))


def cell_range(field: Field) -> tuple[str, str]:
    """``(top_left, bottom_right)`` cell coordinates of a field's merged area."""
    top, bottom = band(field)
    c0, c1 = COL_PX.index(_px(field.x0)) + 1, COL_PX.index(_px(field.x1))
    r0, r1 = ROW_PX.index(_px(top)) + 1, ROW_PX.index(_px(bottom))
    return f'{get_column_letter(c0)}{r0}', f'{get_column_letter(c1)}{r1}'


# Overlay-value key → the top-left cell its value is written into.
FIELD_CELLS: tuple[tuple[str, str], ...] = tuple(
    (cell_range(f)[0], f.key) for f in FIELDS if not f.key.startswith('label:')
)
