"""Document render service — the generic, document-agnostic engine.

Pipeline:  registry spec → context builder → docxtpl fill → .docx bytes
                                                    │ (if pdf)
                                       LibreOffice ─┴─→ .pdf bytes

Mirrors the Boss-dashboard export contract (``apps/export/exports/*`` return raw
bytes; the view wraps them in an ``HttpResponse`` with an attachment header).

PDF conversion shells out to LibreOffice headless, which must be installed on the
server (``LIBREOFFICE_BIN`` setting, or ``soffice``/``libreoffice`` on PATH). When
it is absent, ``render_pdf`` raises ``DocumentRenderError`` with a clear message
rather than producing a broken file. The ``.docx`` path has no such dependency.
"""
from __future__ import annotations

import logging
import shutil
import subprocess
import tempfile
import time
import zipfile
from contextlib import contextmanager
from io import BytesIO
from pathlib import Path

import openpyxl
from django.conf import settings
from django.core.cache import cache
from docx.shared import Emu, Mm
from docxtpl import DocxTemplate, InlineImage
from openpyxl.styles import Font

from apps.contracts.document_templates import registry as tpl_registry
from apps.contracts.services import document_context, document_highlight
from apps.contracts.services.document_context import StampImage

logger = logging.getLogger(__name__)

DOCX_CONTENT_TYPE = (
    'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
)
XLSX_CONTENT_TYPE = (
    'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
)
PDF_CONTENT_TYPE = 'application/pdf'

# Per-scope extractor for the registry out_pattern (download filename) fields.
_FILENAME_FIELDS = {
    tpl_registry.SCOPE_INVOICE: document_context.invoice_filename_fields,
    tpl_registry.SCOPE_SHIPMENT: document_context.cmr_filename_fields,
    tpl_registry.SCOPE_CONTRACT: document_context.contract_filename_fields,
}


class DocumentRenderError(RuntimeError):
    """Raised when a document cannot be rendered (e.g. PDF converter missing)."""


class PdfBusyError(DocumentRenderError):
    """No PDF render slot was free within the wait window.

    Subclasses ``DocumentRenderError`` so the views already map it to HTTP 503
    ("service busy") with no extra handling.
    """


# ════════════════════════════════════════════════
# PDF concurrency limit (worker-starvation guard)
# ════════════════════════════════════════════════
# Each PDF request shells out to LibreOffice and blocks a worker for up to 120s.
# With only 3 uvicorn workers, a handful of concurrent PDF requests would stall
# the whole API. We cap the number of renders running at once across ALL workers
# via slot keys in the shared Redis cache; excess requests wait briefly, then get
# a 503 telling them to retry — the API stays responsive for everyone else.
PDF_MAX_CONCURRENCY = getattr(settings, 'PDF_MAX_CONCURRENCY', 2)
# Safety TTL on a held slot: if a worker dies mid-render the slot self-frees.
# Must exceed the 120s LibreOffice timeout so a live render never loses its slot.
_PDF_SLOT_TTL = 180
# How long a request waits for a free slot before giving up with 503.
_PDF_ACQUIRE_TIMEOUT = 8.0
_PDF_ACQUIRE_POLL = 0.25


@contextmanager
def _pdf_render_slot():
    """Hold one of ``PDF_MAX_CONCURRENCY`` render slots for the duration.

    Uses atomic ``cache.add`` (Redis SET NX) so the slot count is enforced across
    all gunicorn/uvicorn workers, not per-process. Raises :class:`PdfBusyError`
    if no slot frees within ``_PDF_ACQUIRE_TIMEOUT``.
    """
    deadline = time.monotonic() + _PDF_ACQUIRE_TIMEOUT
    held_key = None
    while held_key is None:
        for slot in range(PDF_MAX_CONCURRENCY):
            key = f'pdf:render:slot:{slot}'
            if cache.add(key, 1, _PDF_SLOT_TTL):
                held_key = key
                break
        if held_key is not None:
            break
        if time.monotonic() >= deadline:
            raise PdfBusyError(
                'PDF service is busy (too many conversions at once). '
                'Please try again in a moment.'
            )
        time.sleep(_PDF_ACQUIRE_POLL)
    try:
        yield
    finally:
        cache.delete(held_key)


def _resolve_stamp(tpl: DocxTemplate, value):
    """Turn a StampImage marker into a docxtpl InlineImage; '' if no image.

    Reads the FieldFile's bytes here (render layer does the I/O, builders stay
    pure). A missing/unreadable file degrades to '' rather than failing the doc.
    """
    if not isinstance(value, StampImage):
        return value
    field = value.file
    if not field or not getattr(field, 'name', ''):
        return ''
    try:
        field.open('rb')
        try:
            data = field.read()
        finally:
            field.close()
    except (OSError, ValueError):
        logger.warning('stamp image unreadable: %s', getattr(field, 'name', '?'))
        return ''
    return InlineImage(tpl, BytesIO(data), width=Mm(value.width_mm))


def render_docx(
    template_path: Path, context: dict, highlight: bool = True, layout=None,
) -> bytes:
    """Fill a ``.docx`` template with a Jinja context and return OOXML bytes.

    ``StampImage`` markers in the context are resolved to ``InlineImage`` here
    (they need the ``DocxTemplate``); every other value passes through untouched.

    Args:
        template_path: The ``.docx`` template to fill.
        context: The Jinja context from the registry's builder.
        highlight: Render database-filled values in red (see
            :mod:`~apps.contracts.services.document_highlight`). Stamps are
            resolved BEFORE wrapping so the ``InlineImage`` objects are never
            sentinel-wrapped.
        layout: Optional ``DocumentLayoutSetting`` whose margin/font/spacing
            adjustments are applied after rendering.
    """
    tpl = DocxTemplate(str(template_path))
    context = {key: _resolve_stamp(tpl, value) for key, value in context.items()}
    if highlight:
        context = document_highlight.wrap_context(context)
    tpl.render(context)
    if highlight:
        document_highlight.colorize(tpl.docx)
    if layout is not None and not layout.is_default:
        apply_layout(tpl.docx, layout)
    buf = BytesIO()
    tpl.save(buf)
    return buf.getvalue()


# ════════════════════════════════════════════════
# Page layout adjustments (DocumentLayoutSetting)
# ════════════════════════════════════════════════

def layout_for(document_key: str):
    """The saved layout adjustments for a document key, or ``None``.

    Deliberately NOT cached. The read is one indexed row from a table with at most
    six rows — negligible next to the render it precedes (a PDF shells out to
    LibreOffice for 10-30s). Caching it would buy nothing measurable and would put
    a staleness window in the middle of the one loop that matters: the operator
    nudges a slider, re-downloads, and looks. A stale layout there reads as "the
    setting didn't work" and sends them round again.

    Imported lazily — ``services`` is imported from ``models`` in places, and a
    module-level model import would close the cycle.
    """
    if not tpl_registry.supports_layout(document_key):
        return None
    from apps.contracts.models import DocumentLayoutSetting

    return DocumentLayoutSetting.objects.filter(document_key=document_key).first()


def apply_layout(doc, layout) -> None:
    """Apply a document's saved layout adjustments to a rendered document.

    Order matters: margins move first so the table rescale can measure the change
    it has to compensate for.
    """
    _apply_margins(doc, layout)
    _apply_font_scale(doc, layout.font_scale_pct)
    if layout.line_spacing is not None:
        # Every shipped template leaves line spacing to the Normal style — not one
        # paragraph in any of them sets it explicitly — so this reaches all of them
        # without a per-paragraph walk.
        doc.styles['Normal'].paragraph_format.line_spacing = float(layout.line_spacing)


def _apply_margins(doc, layout) -> None:
    """Add the margin deltas to every section, then rescale pinned tables.

    Deltas rather than absolutes because ``contract_kz.docx`` has two sections with
    deliberately different top margins — an absolute value would flatten them.
    """
    deltas = layout.margin_deltas_mm
    if not any(deltas.values()):
        return

    # The rescale must run exactly ONCE, outside the section loop: `doc.tables`
    # spans the whole document, so rescaling per section would apply the ratio
    # once per section — squaring it on contract_kz's two-section layout.
    # Sections share a page width and take the same deltas, so section 0's ratio
    # governs; the sections' own content widths differ by ~2mm, which moves the
    # ratio by ~0.1%.
    first = doc.sections[0]
    before = first.page_width - first.left_margin - first.right_margin

    for section in doc.sections:
        for attribute, delta_mm in deltas.items():
            current = getattr(section, attribute)
            # Clamp at 0: a negative margin is not a thing Word can print.
            setattr(section, attribute, max(Emu(0), current + Mm(delta_mm)))

    after = first.page_width - first.left_margin - first.right_margin
    if after != before and before:
        _rescale_tables(doc, after / before)


def _rescale_tables(doc, ratio: float) -> None:
    """Scale every fixed-width table by ``ratio`` so it still fits the text area.

    ``build_templates._col_widths`` pins ``autofit = False`` with hard widths sized
    to exactly the current content width; widen a margin and those tables would run
    off the page. Scaling is relative to each table's OWN current total, not an
    assumed page width — ``contract_kz.docx``'s hand-authored tables legitimately
    start wider than their section, and that relationship is not ours to "fix".

    ``autofit = True`` tables are skipped: Word reflows those itself.

    Top-level tables only — unlike ``document_highlight.colorize`` this does not
    recurse into cells. No shipped template nests a table; revisit if one does.
    """
    for table in doc.tables:
        if table.autofit:
            continue
        for column in table.columns:
            if column.width:
                column.width = Emu(int(column.width * ratio))
        for row in table.rows:
            for cell in row.cells:
                if cell.width:
                    cell.width = Emu(int(cell.width * ratio))


# Word stores a font size as whole half-points (``<w:sz w:val>``), so a scaled
# size has to land on that grid. Truncating would silently drop up to half a
# point on every run — at the 8pt sizes these templates use, that is a visible
# step. Round to the nearest half-point instead, and never scale a run away.
_HALF_POINT_EMU = 6350


def _scaled_font_size(size, ratio: float):
    """``size * ratio`` snapped to the nearest half-point Word can store."""
    half_points = max(1, round(size * ratio / _HALF_POINT_EMU))
    return Emu(half_points * _HALF_POINT_EMU)


def _apply_font_scale(doc, scale_pct: int) -> None:
    """Multiply every run's font size by ``scale_pct``.

    Setting ``styles['Normal']`` alone is close to a no-op: most runs carry an
    explicit ``<w:sz>`` that overrides it (713 of 879 runs in ``contract_kz.docx``,
    70 of 83 in ``invoice_ru.docx``). Normal is scaled too, for the minority of
    runs that do inherit.
    """
    if scale_pct == 100:
        return
    ratio = scale_pct / 100

    normal = doc.styles['Normal'].font
    if normal.size:
        normal.size = _scaled_font_size(normal.size, ratio)

    for paragraph in document_highlight.all_paragraphs(doc):
        for run in paragraph.runs:
            if run.font.size:
                run.font.size = _scaled_font_size(run.font.size, ratio)


def render_xlsx(template_path: Path, cell_values: dict, highlight: bool = True) -> bytes:
    """Fill an ``.xlsx`` overlay template by cell coordinate and return OOXML bytes.

    The template's geometry (column widths, row heights, print scale, merges) is
    preserved untouched — the builder returns a ``{coordinate: value}`` map that is
    written into the single active sheet. Used by the CMR print-overlay, whose
    layout must register on the pre-printed official form.

    Every cell written here is by definition a filled value, so ``highlight``
    colours them directly — no sentinels needed as in the ``.docx`` path. The
    template's own font attributes are carried onto the replacement ``Font``;
    openpyxl style objects are shared between cells, so a new one is constructed
    rather than mutated in place.
    """
    wb = openpyxl.load_workbook(template_path)
    ws = wb.active
    for coord, value in cell_values.items():
        cell = ws[coord]
        cell.value = value
        if highlight and value not in (None, ''):
            old = cell.font
            cell.font = Font(
                name=old.name, size=old.size, bold=old.bold, italic=old.italic,
                color=f'FF{document_highlight.HIGHLIGHT_RGB}',
            )
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _libreoffice_bin() -> str | None:
    """Locate the LibreOffice binary (setting first, then PATH)."""
    configured = getattr(settings, 'LIBREOFFICE_BIN', '') or ''
    if configured:
        return configured
    return shutil.which('soffice') or shutil.which('libreoffice')


def render_pdf(source_bytes: bytes, source_ext: str = 'docx') -> bytes:
    """Convert filled ``.docx`` / ``.xlsx`` bytes to PDF via LibreOffice headless.

    A unique ``-env:UserInstallation`` profile per call avoids the shared-profile
    lock that serializes/breaks concurrent headless conversions. LibreOffice picks
    its input filter from the file extension, so the source is written with its
    native ``source_ext`` (``docx`` for Word templates, ``xlsx`` for the CMR overlay).

    Raises:
        DocumentRenderError: If LibreOffice is not installed or conversion fails.
    """
    binary = _libreoffice_bin()
    if not binary:
        raise DocumentRenderError(
            'PDF export requires LibreOffice (set LIBREOFFICE_BIN or install '
            'soffice/libreoffice on the server). The native (.docx/.xlsx) export '
            'is unaffected.'
        )

    with tempfile.TemporaryDirectory() as tmp:
        tmp_path = Path(tmp)
        src = tmp_path / f'in.{source_ext}'
        src.write_bytes(source_bytes)
        profile = (tmp_path / 'lo_profile').as_uri()
        try:
            with _pdf_render_slot():
                subprocess.run(
                    [
                        binary, '--headless', f'-env:UserInstallation={profile}',
                        '--convert-to', 'pdf', '--outdir', str(tmp_path), str(src),
                    ],
                    check=True, capture_output=True, timeout=120,
                )
        except (subprocess.CalledProcessError, subprocess.TimeoutExpired) as exc:
            logger.error('LibreOffice PDF conversion failed: %s', exc, exc_info=True)
            raise DocumentRenderError('PDF conversion failed.') from exc

        pdf = tmp_path / 'in.pdf'
        if not pdf.exists():
            raise DocumentRenderError('PDF conversion produced no output.')
        return pdf.read_bytes()


def generate(
    document_key: str, primary_obj, fmt: str = 'docx', overrides: dict | None = None,
    highlight: bool = True,
) -> tuple[bytes, str, str]:
    """Render a registered document for a primary object.

    Args:
        document_key: A registry key (e.g. ``invoice_ru``).
        primary_obj: The Invoice/Contract/Shipment instance to render from.
        fmt: ``'docx'`` or ``'pdf'``.
        overrides: Optional generate-time field values (e.g. ``place_loading``,
            ``tir_carnet``) passed through to the context builder.
        highlight: Render database-filled values in red. This is a *render* flag,
            not a context value — it never reaches the builders.

    Returns:
        ``(file_bytes, filename_with_extension, content_type)``.

    Raises:
        KeyError: Unknown document_key.
        ValueError: Unsupported fmt.
        DocumentRenderError: PDF requested but converter unavailable / failed.
    """
    # TODO(docs): persist generated documents.
    #   1. GeneratedDocument audit model — one row per generation (document_key,
    #      shipment/sale, user, timestamp) to back the "13:00 board": at a glance
    #      per truck, what's done vs pending. Doesn't block generation; it's
    #      progress tracking for the document team.
    #   2. Save-and-reuse — store the rendered file so that, if nothing about the
    #      truck changed since, re-download serves the saved copy instead of
    #      re-rendering (faster, and the PDF the office printed is preserved).
    #      Needs a "dirty since generated" check (e.g. compare shipment/sale
    #      updated_at vs the saved doc's timestamp) to know when to regenerate.
    if fmt not in ('docx', 'pdf'):
        raise ValueError(f'Unsupported format: {fmt!r}')

    spec = tpl_registry.get_spec(document_key)
    builder = tpl_registry.resolve_builder(spec)
    context = builder(primary_obj, spec.language, overrides)

    if spec.engine == 'xlsx':
        source_bytes = render_xlsx(spec.template_path, context, highlight)
        native_ext, native_type = 'xlsx', XLSX_CONTENT_TYPE
    else:
        source_bytes = render_docx(
            spec.template_path, context, highlight, layout_for(document_key),
        )
        native_ext, native_type = 'docx', DOCX_CONTENT_TYPE

    fields = _FILENAME_FIELDS[spec.scope](primary_obj)
    stem = spec.out_pattern.format(**fields)

    if fmt == 'pdf':
        return render_pdf(source_bytes, native_ext), f'{stem}.pdf', PDF_CONTENT_TYPE
    return source_bytes, f'{stem}.{native_ext}', native_type


ZIP_CONTENT_TYPE = 'application/zip'

# Per-firm request letters bundled into the packet (single-language forms).
_PACKET_LETTER_KEYS = ('ct1_ru', 'fito_ru', 'customs_tk')


def generate_packet_zip(
    shipment, sales, lang='ru', fmt='docx', overrides=None, highlight=True,
) -> bytes:
    """Bundle a truck's whole document packet into a single zip.

    Contents: the truck-level CMR, then per firm (each ``ContractSale`` in
    ``sales``) that firm's invoice + the CT-1 / FITO / customs request letters.
    Filenames are the same per-document names ``generate`` produces, so nothing
    collides (invoice/CMR names carry the contract/invoice/shipment code).

    Args:
        shipment: the ``Shipment`` (truck) — drives the CMR.
        sales: its ``ContractSale`` rows (firms with a linked contract). Firms
            without a sale have no invoice/letters, so they're simply absent.
        lang: ``'ru'`` | ``'en'`` — CMR + invoice language (letters are fixed-lang).
        fmt: ``'docx'`` | ``'pdf'`` — applied to every document in the packet.
        overrides: generate-time values (``place_loading`` / ``tir_carnet``).
        highlight: red-fill flag, applied to every document in the packet.

    Raises:
        ValueError / DocumentRenderError: propagated from ``generate`` (e.g. PDF
        requested but LibreOffice missing) — the whole packet fails as one.
    """
    # The Word CMR (``cmr_*_docx``) is the form the office uses — same as the CMR
    # button; the xlsx overlay (``cmr_*``) is not offered in the UI.
    cmr_key = 'cmr_en_docx' if lang == 'en' else 'cmr_ru_docx'
    invoice_key = 'invoice_en' if lang == 'en' else 'invoice_ru'

    buf = BytesIO()
    with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as archive:
        data, filename, _ = generate(cmr_key, shipment, fmt, overrides, highlight)
        archive.writestr(filename, data)
        for sale in sales:
            for key in (invoice_key, *_PACKET_LETTER_KEYS):
                data, filename, _ = generate(key, sale, fmt, overrides, highlight)
                archive.writestr(filename, data)
    return buf.getvalue()
