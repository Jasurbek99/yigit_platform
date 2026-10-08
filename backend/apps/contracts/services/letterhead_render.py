"""Put a request letter on the seller's letterhead (spec 2026-10-05 §5).

The rendered letter stays the master document, so its fonts, margins and saved
layout survive; the letterhead body is inserted at the top. Once merged, the
letterhead's paragraphs and runs would resolve their styles against the
LETTER's styles (both files have a "Normal"), so before insertion every spacing
and font property the letterhead inherits from its own styles is written onto
its paragraphs and runs directly. Found by the 2026-10-05 prototype: without it
the Yigit blank took the letter's line spacing and pushed ARZA onto a 2nd page.
"""
from __future__ import annotations

import logging
import re
from io import BytesIO

from docx import Document
from docx.enum.section import WD_SECTION
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docxcompose.composer import Composer
from lxml import etree

from apps.core.letterhead import UNREADABLE_DOCX_ERRORS, fill_number

logger = logging.getLogger(__name__)

_SPACING_ATTRS = ('before', 'after', 'line', 'lineRule')
# Word's built-in values when nothing in the chain sets an attribute.
_SPACING_BUILTIN = {'before': '0', 'after': '0', 'line': '240', 'lineRule': 'auto'}
_FONT_ATTRS = ('ascii', 'hAnsi', 'cs', 'eastAsia')


def _style_chain(style) -> list:
    chain = []
    while style is not None:
        chain.append(style.element)
        style = style.base_style
    return chain


def _doc_default(doc, kind: str):
    """docDefaults' ``pPr`` (kind='p') or ``rPr`` (kind='r'), or None."""
    defaults = doc.styles.element.find(qn('w:docDefaults'))
    if defaults is None:
        return None
    tag = 'w:pPrDefault' if kind == 'p' else 'w:rPrDefault'
    return defaults.find(f"{qn(tag)}/{qn('w:' + kind + 'Pr')}")


def _paragraphs(doc):
    """Every paragraph of the body, including ones inside text boxes and tables."""
    return [Paragraph(p, doc._body) for p in doc.element.body.iter(qn('w:p'))]


def _resolve(sources, child: str, attrs, own=None) -> dict:
    """Per-attribute inheritance: first value found in ``own``, then each source's ``child``."""
    values = {a: (own.get(qn(f'w:{a}')) if own is not None else None) for a in attrs}
    for source in sources:
        element = source.find(qn(child)) if source is not None else None
        if element is None:
            continue
        for attr in attrs:
            if values[attr] is None:
                values[attr] = element.get(qn(f'w:{attr}'))
    return values


def _freeze_spacing(doc) -> None:
    """Write every spacing attribute onto each letterhead paragraph explicitly.

    OOXML inherits w:spacing PER ATTRIBUTE, so copying one inherited element is not
    enough — a missing ``line`` would still come from the letter's Normal.
    """
    default_ppr = _doc_default(doc, 'p')
    for paragraph in _paragraphs(doc):
        ppr = paragraph._p.get_or_add_pPr()
        sources = [s.find(qn('w:pPr')) for s in _style_chain(paragraph.style)] + [default_ppr]
        values = _resolve(sources, 'w:spacing', _SPACING_ATTRS, own=ppr.find(qn('w:spacing')))
        spacing = ppr.get_or_add_spacing()
        for attr in _SPACING_ATTRS:
            spacing.set(qn(f'w:{attr}'), values[attr] or _SPACING_BUILTIN[attr])


def _freeze_run_fonts(doc) -> None:
    """Same per-attribute freeze for run font face (w:rFonts) and size (w:sz).

    Only explicit faces are copied: theme-font attributes (asciiTheme …) name the
    letterhead's theme, which does not travel with the merge.
    """
    default_rpr = _doc_default(doc, 'r')
    for paragraph in _paragraphs(doc):
        para_chain = [s.find(qn('w:rPr')) for s in _style_chain(paragraph.style)]
        for run in paragraph.runs:
            rpr = run._r.get_or_add_rPr()
            char_chain = [s.find(qn('w:rPr')) for s in _style_chain(run.style)]
            sources = char_chain + para_chain + [default_rpr]
            fonts = _resolve(sources, 'w:rFonts', _FONT_ATTRS, own=rpr.find(qn('w:rFonts')))
            if any(fonts.values()):
                r_fonts = rpr.get_or_add_rFonts()
                for attr, value in fonts.items():
                    if value:
                        r_fonts.set(qn(f'w:{attr}'), value)
            size = _resolve(sources, 'w:sz', ('val',), own=rpr.find(qn('w:sz')))['val']
            if size:
                rpr.get_or_add_sz().set(qn('w:val'), size)


_A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
_HEX = re.compile(r'^[0-9A-Fa-f]{6}$')
# Names a colour reference may use → the theme's colour-scheme slot (default
# clrSchemeMapping: text = dark, background = light).
_THEME_SLOT = {
    'tx1': 'dk1', 'bg1': 'lt1', 'tx2': 'dk2', 'bg2': 'lt2',
    'text1': 'dk1', 'background1': 'lt1', 'text2': 'dk2', 'background2': 'lt2',
    'dark1': 'dk1', 'light1': 'lt1', 'dark2': 'dk2', 'light2': 'lt2',
    'hyperlink': 'hlink', 'followedHyperlink': 'folHlink',
}
# (theme attribute, its tint/shade companions, the attribute holding the concrete colour)
_W_THEME_ATTRS = (
    ('themeColor', ('themeTint', 'themeShade'), ('val', 'color')),
    ('themeFill', ('themeFillTint', 'themeFillShade'), ('fill',)),
)


def _theme_palette(doc) -> dict:
    """The letterhead theme's colour scheme as ``{slot: 'RRGGBB'}``."""
    theme = next(
        (p for p in doc.part.package.iter_parts() if str(p.partname).startswith('/word/theme/')),
        None,
    )
    if theme is None:
        return {}
    # The theme comes from an uploaded file: parse it with python-docx's hardened
    # parser (no entity resolution), like every other part of the document.
    scheme = parse_xml(theme.blob).find(f'.//{{{_A}}}clrScheme')
    palette = {}
    for slot in (scheme if scheme is not None else []):
        rgb = slot.find(f'{{{_A}}}srgbClr')
        sys_clr = slot.find(f'{{{_A}}}sysClr')
        value = rgb.get('val') if rgb is not None else (sys_clr.get('lastClr') if sys_clr is not None else None)
        if value and _HEX.match(value):
            palette[etree.QName(slot).localname] = value
    return palette


def _freeze_theme_colors(doc) -> None:
    """Replace the letterhead's THEME colour references with their RGB values.

    A Word blank often colours its text and shapes through the theme ("accent6")
    rather than a fixed RGB. Once merged, the reference resolves against the
    LETTER's theme: the Yigit blank's green (70AD47) turned orange (F79646).
    """
    palette = _theme_palette(doc)

    def rgb(name):
        return palette.get(_THEME_SLOT.get(name, name))

    body = doc.element.body
    for element in body.iter():
        for theme_attr, companions, concrete_attrs in _W_THEME_ATTRS:
            name = element.get(qn(f'w:{theme_attr}'))
            if name is None:
                continue
            concrete = next((a for a in concrete_attrs if element.get(qn(f'w:{a}')) is not None), concrete_attrs[0])
            # Word stores the resolved colour (tint/shade applied) next to the
            # reference; use it, falling back to the plain theme colour.
            if not _HEX.match(element.get(qn(f'w:{concrete}')) or '') and rgb(name):
                element.set(qn(f'w:{concrete}'), rgb(name))
            for attr in (theme_attr, *companions):
                element.attrib.pop(qn(f'w:{attr}'), None)
    for scheme in list(body.iter(f'{{{_A}}}schemeClr')):
        value = rgb(scheme.get('val'))
        if value is None:  # e.g. phClr — a placeholder the shape style fills in
            continue
        fixed = etree.Element(f'{{{_A}}}srgbClr', val=value)
        fixed.extend(list(scheme))  # keep lumMod / tint / shade modifiers
        scheme.getparent().replace(scheme, fixed)


def _trim_trailing_empty_paragraphs(doc) -> None:
    """Drop the letterhead's empty paragraphs after its last content.

    A Word blank ends in a few empty lines left for typing; once the letter sits
    under it those only push the letter down — the Yigit blank's three pushed
    ARZA onto a 2nd page. Stops at a paragraph with text, a drawing or a sectPr,
    and never removes the paragraph right after a table (Word needs one there).
    """
    body = doc.element.body
    while True:
        children = [e for e in body if e.tag != qn('w:sectPr')]
        if len(children) < 2 or children[-1].tag != qn('w:p'):
            return
        last, before = children[-1], children[-2]
        if (''.join(last.itertext()).strip()
                or last.find('.//' + qn('w:sectPr')) is not None
                or last.find('.//' + qn('w:drawing')) is not None
                or before.tag == qn('w:tbl')):
            return
        body.remove(last)


def apply_letterhead(letter_bytes: bytes, letterhead_bytes: bytes, number: int | None) -> bytes:
    """The letter with the letterhead body on top and ``number`` in its «№ ___»."""
    head = Document(BytesIO(letterhead_bytes))
    fill_number(head, number)
    _freeze_spacing(head)
    _freeze_run_fonts(head)
    _freeze_theme_colors(head)
    _trim_trailing_empty_paragraphs(head)
    letter = Document(BytesIO(letter_bytes))
    # Sections of different page sizes cannot share a page, so a US Letter blank
    # would push an A4 letter onto page 2: the blank takes the letter's size.
    page = letter.sections[-1]
    for section in head.sections:
        section.page_width, section.page_height = page.page_width, page.page_height
    composer = Composer(letter)
    composer.insert(0, head)
    merged = composer.doc
    if len(merged.sections) > 1:
        # The letter follows the letterhead's section break on the SAME page.
        merged.sections[-1].start_type = WD_SECTION.CONTINUOUS
    buf = BytesIO()
    merged.save(buf)
    return buf.getvalue()


def letterhead_for(sale) -> bytes | None:
    """The seller firm's letterhead bytes, or None (no file / unreadable → plain letter)."""
    contract = sale.contract
    firm = sale.export_firm or (contract.export_firm if contract else None)
    field = getattr(firm, 'letterhead', None)
    if not field or not getattr(field, 'name', ''):
        return None
    try:
        field.open('rb')
        try:
            data = field.read()
        finally:
            field.close()
        Document(BytesIO(data))
    except (OSError, *UNREADABLE_DOCX_ERRORS):
        logger.warning('letterhead unreadable for export firm %s: %s', firm.pk, field.name)
        return None
    return data
